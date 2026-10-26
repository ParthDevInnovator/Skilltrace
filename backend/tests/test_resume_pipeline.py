"""Offline resume ingestion regression tests with isolated candidate stores."""
import io
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter

from backend.main import app
from backend.models import ClaimSection
from backend.routers.ingest import get_ingestion_store, get_context_builder
from backend.services.context_builder import ContextBuilder
from backend.routers import candidates as candidate_routes
from backend.services.claim_extractor import extract_claims
from backend.services.deduplicator import ResumeIngestionStore
from backend.services.identity_extractor import ExtractedIdentity, extract_identity
from backend.services.resume_parser import MAX_UPLOAD_BYTES, ParsedResume, parse_resume
from backend.services.seed_loader import SeedLoader

DATA_DIR = Path(__file__).resolve().parents[1] / "data"
FIXTURES = Path(__file__).resolve().parents[2] / "frontend/e2e/fixtures"
TEXT = """Name: Morgan Test
Email: morgan.test@example.com
Title: Senior Python Engineer
Location: Pune
8.5 years of experience
https://github.com/morgan-test
https://www.linkedin.com/in/morgan-test/
https://leetcode.com/u/morgan-test/
Experience
- Built reliable Python services for customers.
- Designed Kubernetes deployment automation.
Projects
- Optimized PostgreSQL query performance by 30 percent.
Education
- Implemented a Python compiler project.
"""


@pytest.fixture
def store():
    repository = SeedLoader(data_dir=DATA_DIR)
    repository.load_all()
    return ResumeIngestionStore(repository)


@pytest.fixture
async def api_store(store, monkeypatch):
    async def no_signals(*args, **kwargs):
        return []
    service = ContextBuilder(store, fetchers={name: no_signals for name in ("github", "leetcode", "linkedin")})
    store.context_builder = service
    app.dependency_overrides[get_ingestion_store] = lambda: store
    app.dependency_overrides[get_context_builder] = lambda: service
    monkeypatch.setattr(candidate_routes, "seed_loader", store.repository)
    yield store
    await service.shutdown()
    app.dependency_overrides.pop(get_ingestion_store, None)
    app.dependency_overrides.pop(get_context_builder, None)


def test_parse_documents_and_encoding():
    assert "Build reliable Python" in parse_resume(raw_bytes=(FIXTURES / "role.pdf").read_bytes(), filename="role.pdf").raw_text
    document = Document()
    for line in TEXT.splitlines():
        document.add_paragraph(line)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table skills: FastAPI"
    buffer = io.BytesIO(); document.save(buffer)
    parsed = parse_resume(raw_bytes=buffer.getvalue(), filename="resume.docx")
    assert parsed.source == "docx" and "Table skills: FastAPI" in parsed.raw_text
    utf16 = parse_resume(raw_bytes=TEXT.encode("utf-16"), filename="resume.TXT")
    assert utf16.source == "txt" and "Morgan Test" in utf16.raw_text
    assert utf16.char_count == len(utf16.raw_text)
    assert parse_resume(path=FIXTURES / "role.txt").source == "txt"


def test_seed_fallback_and_failure_modes():
    parsed = parse_resume(path="missing-resume.pdf", stub_text=TEXT)
    assert parsed.source == "stub" and parsed.raw_text == TEXT.strip()
    assert parse_resume(stub_text=TEXT).source == "stub"
    assert parse_resume().source == "empty"
    assert parse_resume(path="missing-resume.pdf").source == "empty"
    assert parse_resume(raw_bytes=b"not a PDF", filename="bad.pdf", stub_text=TEXT).source == "empty"
    assert parse_resume(raw_bytes=b"broken zip", filename="bad.docx").source == "empty"
    assert parse_resume(raw_bytes=b"", filename="empty.txt").source == "empty"
    assert parse_resume(raw_bytes=b"valid text", filename="bad.exe").source == "empty"


def blank_pdf():
    writer = PdfWriter(); writer.add_blank_page(width=300, height=300)
    output = io.BytesIO(); writer.write(output)
    return output.getvalue()


def test_scanned_pdf_is_detected_without_raising():
    parsed = parse_resume(raw_bytes=blank_pdf(), filename="scan.pdf")
    assert parsed.source == "pdf" and parsed.char_count == 0
    assert parsed.warning == "insufficient_pdf_text"


def test_identity_extraction_is_nullable_and_platform_specific():
    identity = extract_identity(TEXT)
    assert identity.full_name == "Morgan Test"
    assert identity.email == "morgan.test@example.com"
    assert identity.location == "Pune"
    assert identity.headline == "Senior Python Engineer"
    assert identity.experience_years == 8.5
    assert identity.github == "https://github.com/morgan-test"
    assert identity.linkedin == "https://linkedin.com/in/morgan-test"
    assert identity.leetcode == "https://leetcode.com/u/morgan-test"
    assert all(value is None for value in extract_identity("").model_dump().values())
    assert extract_identity("Morgan Test\nPython Engineer").full_name == "Morgan Test"
    assert extract_identity("https://evil.github.com/morgan-test").github is None
    assert extract_identity("https://github.com/morgan-test/repo").github is None


def test_claims_have_sections_triggers_stable_ids_and_no_verification():
    claims = extract_claims(TEXT, "CAND-001")
    assert len(claims) == 4
    assert {claim.extraction_trigger for claim in claims} >= {"Built", "Designed", "Optimized"}
    assert [claim.section for claim in claims] == [ClaimSection.EXPERIENCE, ClaimSection.EXPERIENCE, ClaimSection.PROJECTS, ClaimSection.EDUCATION]
    assert all(claim.verification_state == "unverified" for claim in claims)
    assert [claim.id for claim in claims] == [claim.id for claim in extract_claims(TEXT, "CAND-001")]
    assert len(extract_claims(TEXT + "\n- BUILT   reliable Python services for customers.", "CAND-001")) == 4
    assert extract_claims("- Built a tool without a recognized skill", "CAND-001") == []
    assert extract_claims("Skills\n- Python", "CAND-001") == []
    assert extract_claims("- Shipped Erlang services", "CAND-001", action_verbs=["Shipped"], skill_keywords=["Erlang"])[0].skill == "Erlang"
    assert len(extract_claims("\n".join(f"- Built Python service number {i}" for i in range(30)), "CAND-001")) == 20


@pytest.mark.asyncio
async def test_email_dedup_and_updates_preserve_missing_metadata(store):
    identity = extract_identity(TEXT)
    parsed = ParsedResume(raw_text=TEXT, source="txt")
    first, updated, revision = store.upsert(identity, parsed, "resume.txt", TEXT.encode())
    assert first.id == "CAND-007" and not updated
    second, updated, newer = store.upsert(ExtractedIdentity(email="MORGAN.TEST@example.com"), parsed, "new.txt", TEXT.encode())
    assert second.id == first.id and updated and newer != revision
    assert second.full_name == first.full_name and second.location == first.location
    async def no_signals(*args, **kwargs):
        return []
    service = ContextBuilder(store, fetchers={source: no_signals for source in ("github", "leetcode", "linkedin")})
    assert service.reserve(first.id, revision) is None
    assert store.pending_builds[first.id].stage == "queued"  # Stale upload cannot publish.
    service.launch(first.id, service.reserve(first.id, newer))
    await service.drain()
    assert store.pending_builds[first.id].stage == "ready"
    assert store.pending_builds[first.id].status == "ready"
    assert len(store.pending_builds[first.id].claims) == 4
    assert len(store.repository.get_candidate(first.id).claims_preview) == 4


def test_parallel_uploads_are_atomic(store):
    parsed = ParsedResume(raw_text=TEXT, source="txt")
    def upload(index):
        identity = ExtractedIdentity(full_name="Morgan Test", email="shared@example.com")
        return store.upsert(identity, parsed, "resume.txt", TEXT.encode())[0].id
    with ThreadPoolExecutor(max_workers=8) as executor:
        ids = list(executor.map(upload, range(16)))
    assert set(ids) == {"CAND-007"}
    def distinct(index):
        return store.upsert(ExtractedIdentity(full_name="Different Candidate", email=f"distinct{index}@example.com"), parsed, "resume.txt", TEXT.encode())[0].id
    with ThreadPoolExecutor(max_workers=8) as executor:
        assert len(set(executor.map(distinct, range(16)))) == 16


@pytest.mark.asyncio
async def test_upload_creates_and_updates_candidate_and_stages_context(api_store):
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/api/v1/ingest/resume", files={"file": ("resume.txt", TEXT.encode(), "text/plain")})
        assert response.status_code == 202, response.text
        result = response.json()
        assert result["candidate_id"] == "CAND-007" and result["status"] == "building"
        assert result["parse_source"] == "txt" and result["enrichment_pending"]
        await api_store.context_builder.drain()
        pending = api_store.pending_builds[result["candidate_id"]]
        assert pending.stage == "ready" and pending.original_bytes == TEXT.encode()
        assert api_store.context_builder.registry.get(result["candidate_id"]).status == "ready"
        assert pending.parsed_resume.source == "txt"
        detail = await client.get(f"/api/v1/candidates/{result['candidate_id']}")
        assert detail.status_code == 200
        assert detail.json()["full_name"] == "Morgan Test"
        listed = await client.get("/api/v1/candidates", params={"search": "morgan.test@example.com"})
        assert [candidate["id"] for candidate in listed.json()] == [result["candidate_id"]]
        response = await client.post("/api/v1/ingest/resume", files={"file": ("resume.txt", TEXT.encode())})
        assert response.status_code == 202 and response.json()["updated_existing"]
        assert api_store.repository.stats()["total_candidates"] == 7


@pytest.mark.asyncio
@pytest.mark.parametrize("filename,data,expected", [
    ("resume.exe", b"invalid", 422), ("empty.txt", b"", 422),
    ("huge.txt", b"a" * (MAX_UPLOAD_BYTES + 1), 413),
    ("corrupt.pdf", b"invalid", 422), ("scan.pdf", blank_pdf(), 422),
    ("identity.txt", b"A resume with no identity.", 422),
], ids=["unsupported", "empty", "oversized", "corrupt", "scanned", "missing-identity"])
async def test_upload_rejections_do_not_mutate_store(api_store, filename, data, expected):
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/api/v1/ingest/resume", files={"file": (filename, data)})
        assert response.status_code == expected, response.text
        assert api_store.repository.stats()["total_candidates"] == 6
        assert api_store.pending_builds == {}


@pytest.mark.asyncio
async def test_partial_identity_can_be_corrected_with_form_fields(api_store):
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/api/v1/ingest/resume", files={"file": ("role.pdf", (FIXTURES / "role.pdf").read_bytes())},
                                     data={"full_name": "New Candidate", "email": "new@example.com", "location": "London"})
        assert response.status_code == 202, response.text
        candidate = api_store.repository.get_candidate(response.json()["candidate_id"])
        assert candidate.location == "London" and candidate.headline == "Not provided"
        assert candidate.resume_text_stub is None and candidate.match_preview is None
        response = await client.post("/api/v1/ingest/resume", files={"file": ("resume.txt", TEXT.encode())}, data={"email": "invalid"})
        assert response.status_code == 422


@pytest.mark.asyncio
async def test_docx_upload_and_existing_seed_update(api_store):
    document = Document()
    document.add_paragraph("Email: ALEX.RIVERA@example.com")
    document.add_paragraph("Experience")
    document.add_paragraph("- Built Python services with FastAPI.")
    stream = io.BytesIO()
    document.save(stream)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/api/v1/ingest/resume", files={"file": ("update.docx", stream.getvalue())})
    assert response.status_code == 202, response.text
    assert response.json()["candidate_id"] == "CAND-001"
    assert response.json()["updated_existing"] and response.json()["parse_source"] == "docx"
    candidate = api_store.repository.get_candidate("CAND-001")
    assert candidate.full_name == "Alex Rivera"
    assert candidate.external_identifiers.github == "https://github.com/alexrivera-eng"
    assert candidate.resume_path is None and candidate.resume_text_stub is None
    assert candidate.match_preview is None
