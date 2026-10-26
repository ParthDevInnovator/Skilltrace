"""Local document text extraction; no network or candidate inference."""
import io
import logging
import zipfile
from pathlib import Path
from typing import Literal

import chardet
from docx import Document
from pydantic import BaseModel, computed_field
from pypdf import PdfReader

logger = logging.getLogger(__name__)
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ParsedResume(BaseModel):
    raw_text: str = ""
    source: Literal["pdf", "docx", "txt", "stub", "empty"] = "empty"
    warning: str | None = None

    @computed_field
    @property
    def char_count(self) -> int:
        return len(self.raw_text)


def parse_resume(path: str | Path | None = None, stub_text: str | None = None,
                 *, raw_bytes: bytes | None = None, filename: str | None = None) -> ParsedResume:
    """Missing seed files may use a labeled stub; invalid uploads never do."""
    try:
        if raw_bytes is None:
            if path is None or not Path(path).is_file():
                if stub_text and stub_text.strip():
                    return ParsedResume(raw_text=stub_text.strip(), source="stub")
                return ParsedResume(warning="missing_document")
            file_path = Path(path)
            filename = file_path.name
            if file_path.stat().st_size > MAX_UPLOAD_BYTES:
                return ParsedResume(warning="file_too_large")
            with file_path.open("rb") as stream:
                raw_bytes = stream.read(MAX_UPLOAD_BYTES + 1)
        if not raw_bytes:
            return ParsedResume(warning="empty_document")
        if len(raw_bytes) > MAX_UPLOAD_BYTES:
            return ParsedResume(warning="file_too_large")
        extension = Path(filename or "").suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            return ParsedResume(warning="unsupported_format")
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(raw_bytes))
            if reader.is_encrypted or len(reader.pages) > 200:
                return ParsedResume(warning="encrypted_or_large_pdf")
            parts = []
            length = 0
            for page in reader.pages:
                text = page.extract_text() or ""
                length += len(text)
                if length > MAX_TEXT_CHARS:
                    return ParsedResume(warning="text_too_long")
                parts.append(text)
            text = "\n".join(parts)
        elif extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    return ParsedResume(warning="expanded_document_too_large")
            document = Document(io.BytesIO(raw_bytes))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(" ".join(cell.text for cell in row.cells)
                         for table in document.tables for row in table.rows)
            text = "\n".join(parts)
        else:
            encoding = chardet.detect(raw_bytes).get("encoding") or "utf-8"
            text = raw_bytes.decode(encoding)
            if "\x00" in text:
                return ParsedResume(warning="binary_document")
        text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        if len(text) > MAX_TEXT_CHARS:
            return ParsedResume(warning="text_too_long")
        if extension == ".pdf" and len(text) < 50:
            return ParsedResume(raw_text=text, source="pdf", warning="insufficient_pdf_text")
        if not text:
            return ParsedResume(warning="empty_document")
        return ParsedResume(raw_text=text, source=extension[1:])
    except Exception:
        logger.warning("Resume extraction failed; returning an empty parse result")
        return ParsedResume(warning="parse_failed")
